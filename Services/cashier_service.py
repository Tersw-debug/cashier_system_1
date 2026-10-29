"""Cashier operations for MyShop ERP (MySQL 8, migrations 1-6).

Public functions and return shapes are kept compatible with the existing UI.
Money is stored in integer minor currency units. Each business operation writes
its database changes in one transaction; Excel export happens after commit.
"""
from __future__ import annotations

import json
import os
import tempfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from Database.Schema import db

PROJECT_ROOT = Path(__file__).resolve().parents[1]
CASHIER_REPORTS_ROOT = PROJECT_ROOT / "Cashier_Reports"

# ---------------------------------------------------------------------------
# Shared validation, conversion, SQL and audit helpers
# ---------------------------------------------------------------------------
def _money_round(value: Decimal) -> int:
    return int(value.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _decimal_value(value: Any, field_name: str = "القيمة") -> Decimal:
    try:
        result = Decimal(str(value).strip())
    except (InvalidOperation, AttributeError, TypeError, ValueError):
        raise ValueError(f"{field_name} غير صحيحة") from None
    if not result.is_finite():
        raise ValueError(f"{field_name} غير صحيحة")
    return result


def _setting(c, key: str, default: str | None = None) -> str | None:
    row = c.execute(
        "SELECT value FROM settings WHERE `key` = %s LIMIT 1", (key,)
    ).fetchone()
    return row[0] if row else default


def get_currency_decimals() -> int:
    row = db.fetch_one(
        "SELECT value FROM settings WHERE `key` = 'currency_decimals' LIMIT 1"
    )
    try:
        decimals = int(row[0]) if row and row[0] is not None else 2
    except (TypeError, ValueError):
        decimals = 2
    return max(0, min(4, decimals))


def _decimals_or_default(decimals: int | None) -> int:
    if decimals is None:
        return get_currency_decimals()
    try:
        return max(0, min(4, int(decimals)))
    except (TypeError, ValueError, OverflowError):
        raise ValueError("عدد المنازل العشرية غير صحيح") from None


def money_to_minor(value: str | int | float | Decimal, decimals: int | None = None) -> int:
    decimals = _decimals_or_default(decimals)
    try:
        amount = Decimal(str(value).strip().replace(",", ""))
    except (InvalidOperation, AttributeError, TypeError, ValueError):
        raise ValueError("أدخل مبلغًا رقميًا صحيحًا") from None
    if not amount.is_finite() or amount < 0:
        raise ValueError("المبلغ يجب أن يكون رقمًا غير سالب")
    return _money_round(amount * (Decimal(10) ** decimals))


def minor_to_money(value: int | float | Decimal, decimals: int | None = None) -> str:
    decimals = _decimals_or_default(decimals)
    amount = Decimal(str(value)) / (Decimal(10) ** decimals)
    return f"{amount:.{decimals}f}"


def _next_number(c, sequence_name: str) -> str:
    """Allocate a document number safely inside the caller's transaction."""
    row = c.execute(
        """SELECT prefix, next_number, padding
           FROM doc_sequences WHERE name = %s FOR UPDATE""",
        (sequence_name,),
    ).fetchone()
    if row is None:
        raise ValueError(f"تسلسل الترقيم غير موجود: {sequence_name}")

    prefix, number, padding = row
    number, padding = int(number), int(padding)
    if number < 1:
        raise ValueError(f"الرقم التالي غير صالح في التسلسل: {sequence_name}")
    if not 1 <= padding <= 20:
        raise ValueError(f"عدد خانات الترقيم غير صالح في التسلسل: {sequence_name}")

    c.execute(
        "UPDATE doc_sequences SET next_number = %s WHERE name = %s",
        (number + 1, sequence_name),
    )
    return f"{prefix}{number:0{padding}d}"


def _configured_account_id(c, key: str) -> int:
    code = _setting(c, key)
    if code in (None, ""):
        raise ValueError(f"إعداد الحساب غير موجود: {key}")
    row = c.execute(
        "SELECT id FROM accounts WHERE code = %s AND is_active = 1 LIMIT 1",
        (str(code),),
    ).fetchone()
    if not row:
        raise ValueError(f"الحساب {code} غير موجود أو غير نشط")
    return int(row[0])


def _audit(
    c,
    user_id: int | None,
    action: str,
    entity: str,
    entity_id: int,
    new_values: dict[str, Any] | None = None,
    note: str | None = None,
) -> None:
    payload = (
        json.dumps(new_values, ensure_ascii=False, default=str, sort_keys=True)
        if new_values is not None else None
    )
    c.execute(
        """INSERT INTO audit_log
           (user_id, action, entity, entity_id, new_values, note)
           VALUES (%s, %s, %s, %s, %s, %s)""",
        (user_id, action, entity, entity_id, payload, note),
    )


def _require_active_user(c, user_id: int, message: str = "حساب المستخدم غير موجود أو غير نشط") -> None:
    row = c.execute(
        "SELECT is_active FROM users WHERE id = %s FOR UPDATE", (int(user_id),)
    ).fetchone()
    if not row or not row[0]:
        raise ValueError(message)


def _require_open_shift(c, shift_id: int, user_id: int):
    row = c.execute(
        """SELECT user_id, status FROM cash_shifts
           WHERE id = %s FOR UPDATE""",
        (int(shift_id),),
    ).fetchone()
    if not row or int(row[0]) != int(user_id) or row[1] != "OPEN":
        raise ValueError("الوردية المحددة ليست مفتوحة لهذا الكاشير")
    return row


def _post_journal(
    c,
    description: str,
    source_type: str,
    source_id: int,
    user_id: int | None,
    lines: list[dict[str, Any]],
) -> int:
    """Create and post a balanced immutable journal entry in the caller's transaction."""
    if len(lines) < 2:
        raise ValueError("القيد المحاسبي يجب أن يحتوي على بندين على الأقل")

    normalized: list[dict[str, Any]] = []
    for line in lines:
        debit = int(line.get("debit", 0) or 0)
        credit = int(line.get("credit", 0) or 0)
        if debit < 0 or credit < 0 or (debit > 0 and credit > 0) or (debit == 0 and credit == 0):
            raise ValueError("كل بند قيد يجب أن يكون مدينًا أو دائنًا فقط وبقيمة موجبة")
        account_id = int(line["account_id"])
        normalized.append({
            **line,
            "account_id": account_id,
            "debit": debit,
            "credit": credit,
        })

    total_debit = sum(line["debit"] for line in normalized)
    total_credit = sum(line["credit"] for line in normalized)
    if total_debit <= 0 or total_debit != total_credit:
        raise ValueError(f"القيد غير متوازن: المدين {total_debit}، الدائن {total_credit}")

    number = _next_number(c, "JOURNAL")
    c.execute(
        """INSERT INTO journal_entries
           (number, entry_date, description, source_type, source_id, status, user_id)
           VALUES (%s, %s, %s, %s, %s, 'DRAFT', %s)""",
        (number, date.today().isoformat(), description, source_type, int(source_id), user_id),
    )
    entry_id = int(c.lastrowid)

    for line in normalized:
        c.execute(
            """INSERT INTO journal_lines
               (entry_id, account_id, partner_id, debit, credit, description)
               VALUES (%s, %s, %s, %s, %s, %s)""",
            (
                entry_id, line["account_id"], line.get("partner_id"),
                line["debit"], line["credit"], line.get("description") or None,
            ),
        )

    c.execute(
        "UPDATE journal_entries SET status = 'POSTED' WHERE id = %s",
        (entry_id,),
    )
    return entry_id


# ---------------------------------------------------------------------------
# Cashier context, product lookup and customer lookup
# ---------------------------------------------------------------------------
def get_cashier_context(user_id: int) -> dict[str, Any]:
    return {
        "warehouses": db.fetch_all(
            """SELECT id, name, is_default FROM warehouses
               WHERE is_active = 1 ORDER BY is_default DESC, name"""
        ),
        "customers": db.fetch_all(
            """SELECT id, name, COALESCE(phone, '') FROM partners
               WHERE is_customer = 1 AND is_active = 1 ORDER BY name"""
        ),
        "payment_methods": db.fetch_all(
            """SELECT pm.id, pm.name, pm.kind
               FROM payment_methods pm
               JOIN accounts a ON a.id = pm.account_id AND a.is_active = 1
               WHERE pm.is_active = 1 ORDER BY pm.name, pm.id"""
        ),
        "open_shift": get_open_shift(user_id),
        "last_shift": db.fetch_one(
            """SELECT id, opening_cash, opened_at FROM cash_shifts
               WHERE user_id = %s ORDER BY id DESC LIMIT 1""",
            (int(user_id),),
        ),
        "currency_decimals": get_currency_decimals(),
    }



def get_cashier_products(
    query: str = "",
    warehouse_id: int | None = None,
) -> list[tuple]:
    """Search active products and include their complete section hierarchy.

    Existing column positions remain unchanged:
    0=id, 1=name, 2=sku, 3=barcode, 4=sell_price,
    5=stock, 6=min_quantity, 7=track_stock, 8=cost_price,
    9=tax_rate_bp, 10=direct_section_name, 11=color,
    12=complete_section_path.
    """
    query = (query or "").strip()
    params: list[Any] = []

    if warehouse_id is None:
        stock_expr = """
            COALESCE(
                (SELECT SUM(sl.quantity)
                 FROM stock_levels AS sl
                 WHERE sl.product_id = p.id),
                0
            )
        """
    else:
        stock_expr = """
            COALESCE(
                (SELECT sl.quantity
                 FROM stock_levels AS sl
                 WHERE sl.product_id = p.id
                   AND sl.warehouse_id = %s),
                0
            )
        """
        params.append(int(warehouse_id))

    sql = f"""
        WITH RECURSIVE section_paths AS (
            SELECT
                s.id,
                s.parent_id,
                CAST(s.name AS CHAR(1000)) AS full_path
            FROM sections AS s
            WHERE s.parent_id IS NULL

            UNION ALL

            SELECT
                child.id,
                child.parent_id,
                CONCAT(parent.full_path, ' / ', child.name)
            FROM sections AS child
            JOIN section_paths AS parent
              ON child.parent_id = parent.id
        )
        SELECT
            p.id,
            p.name,
            COALESCE(p.sku, ''),
            COALESCE((
                SELECT pb.barcode
                FROM product_barcodes AS pb
                WHERE pb.product_id = p.id
                ORDER BY pb.is_primary DESC, pb.id
                LIMIT 1
            ), ''),
            p.sell_price,
            {stock_expr},
            p.min_quantity,
            p.track_stock,
            p.cost_price,
            COALESCE((
                SELECT tr.rate_bp
                FROM tax_rates AS tr
                WHERE tr.id = p.tax_rate_id
                  AND tr.is_active = 1
            ), 0),
            COALESCE(s.name, ''),
            COALESCE(p.color, ''),
            COALESCE(sp.full_path, s.name, '')
        FROM products AS p
        LEFT JOIN sections AS s
          ON s.id = p.section_id
        LEFT JOIN section_paths AS sp
          ON sp.id = p.section_id
        WHERE p.is_active = 1
    """

    if query:
        sql += """
            AND (
                p.name LIKE %s
                OR p.sku = %s
                OR EXISTS (
                    SELECT 1
                    FROM product_barcodes AS pb
                    WHERE pb.product_id = p.id
                      AND pb.barcode = %s
                )
            )
            ORDER BY
                CASE
                    WHEN p.sku = %s THEN 0
                    WHEN EXISTS (
                        SELECT 1
                        FROM product_barcodes AS pb
                        WHERE pb.product_id = p.id
                          AND pb.barcode = %s
                    ) THEN 0
                    ELSE 1
                END,
                p.name
            LIMIT 100
        """
        params.extend([
            f"%{query}%",
            query,
            query,
            query,
            query,
        ])
    else:
        sql += " ORDER BY p.name LIMIT 500"

    return db.fetch_all(sql, tuple(params))




def get_cashier_product_groups(warehouse_id: int | None = None):
    """Group products by category, size and printing face."""
    rows = get_cashier_products("", warehouse_id)

    grouped: dict[str, list[tuple]] = {}
    ungrouped: list[tuple] = []

    for row in rows:
        # Column 12 contains the full hierarchy; column 10 remains
        # the direct parent name for backwards compatibility.
        path = str(row[12] or "").strip() if len(row) > 12 else ""
        parts = [part.strip() for part in path.split(" / ") if part.strip()]

        if not parts:
            ungrouped.append(row)
            continue

        # Example: جملة / A4 / one face / colorful
        # Group under: جملة / A4 / one face
        group_name = " / ".join(parts[:3])
        grouped.setdefault(group_name, []).append(row)

    return (
        sorted(grouped.items(), key=lambda item: item[0].casefold()),
        ungrouped,
    )



def get_customer_by_name_phone(name: str, phone: str | None = None) -> int | None:
    name = (name or "").strip()
    phone = (phone or "").strip()
    if not name:
        return None
    if phone:
        row = db.fetch_one(
            """SELECT id FROM partners
               WHERE is_customer = 1 AND LOWER(name) = LOWER(%s)
                 AND COALESCE(phone, '') = %s LIMIT 1""",
            (name, phone),
        )
    else:
        row = db.fetch_one(
            """SELECT id FROM partners
               WHERE is_customer = 1 AND LOWER(name) = LOWER(%s) LIMIT 1""",
            (name,),
        )
    return int(row[0]) if row else None


def get_username_id(username: str) -> int | None:
    """Return an active user's ID (fixes the old malformed userLOWER(name) query)."""
    username = str(username or "").strip()
    if not username:
        return None
    row = db.fetch_one(
        """SELECT id FROM users
           WHERE LOWER(username) = LOWER(%s) AND is_active = 1 LIMIT 1""",
        (username,),
    )
    return int(row[0]) if row else None


def get_open_shift(user_id: int):
    return db.fetch_one(
        """SELECT id, opening_cash, opened_at FROM cash_shifts
           WHERE user_id = %s AND status = 'OPEN' ORDER BY id DESC LIMIT 1""",
        (int(user_id),),
    )


def _session_date_from_value(opened_at: Any) -> str:
    if opened_at:
        text = str(opened_at).strip()
        try:
            return datetime.fromisoformat(text.replace("Z", "+00:00")).date().isoformat()
        except ValueError:
            try:
                return date.fromisoformat(text[:10]).isoformat()
            except ValueError:
                pass
    return date.today().isoformat()


def get_cashier_session_folder(shift_id: int, create: bool = True) -> str:
    shift = db.fetch_one(
        "SELECT opened_at FROM cash_shifts WHERE id = %s", (int(shift_id),)
    )
    if not shift:
        raise ValueError("الوردية غير موجودة")
    folder = CASHIER_REPORTS_ROOT / _session_date_from_value(shift[0])
    if create:
        folder.mkdir(parents=True, exist_ok=True)
    return str(folder)


# ---------------------------------------------------------------------------
# Shift open / close
# ---------------------------------------------------------------------------
def open_cashier_shift(user_id: int, opening_cash_minor: int = 0) -> int:
    try:
        opening_cash_minor = int(opening_cash_minor)
    except (TypeError, ValueError, OverflowError):
        raise ValueError("رصيد بداية الوردية غير صحيح") from None
    if opening_cash_minor < 0:
        raise ValueError("رصيد بداية الوردية لا يمكن أن يكون سالبًا")

    # Prepare the likely folder before committing the shift. Resolve it again
    # from the persisted timestamp after commit in case midnight was crossed.
    (CASHIER_REPORTS_ROOT / date.today().isoformat()).mkdir(parents=True, exist_ok=True)
    with db.transaction() as c:
        _require_active_user(c, user_id)
        # Serializes open-shift creation for one user even if two windows race.
        existing = c.execute(
            """SELECT id FROM cash_shifts
               WHERE user_id = %s AND status = 'OPEN' LIMIT 1 FOR UPDATE""",
            (int(user_id),),
        ).fetchone()
        if existing:
            raise ValueError("لديك وردية مفتوحة بالفعل")
        c.execute(
            """INSERT INTO cash_shifts(user_id, opening_cash, note, status)
               VALUES (%s, %s, %s, 'OPEN')""",
            (int(user_id), opening_cash_minor, "بدأت من صفحة الكاشير"),
        )
        shift_id = int(c.lastrowid)
        _audit(c, user_id, "OPEN_SHIFT", "cash_shifts", shift_id,
               {"opening_cash": opening_cash_minor, "status": "OPEN"})

    get_cashier_session_folder(shift_id, create=True)
    return shift_id


def close_cashier_shift(shift_id: int, closing_cash_minor: int, user_id: int) -> int:
    try:
        closing_cash_minor = int(closing_cash_minor)
    except (TypeError, ValueError, OverflowError):
        raise ValueError("النقد الفعلي غير صحيح") from None
    if closing_cash_minor < 0:
        raise ValueError("النقد الفعلي لا يمكن أن يكون سالبًا")

    with db.transaction() as c:
        shift = c.execute(
            """SELECT opening_cash, status FROM cash_shifts
               WHERE id = %s AND user_id = %s FOR UPDATE""",
            (int(shift_id), int(user_id)),
        ).fetchone()
        if not shift or shift[1] != "OPEN":
            raise ValueError("الوردية غير موجودة أو مغلقة")

        cash_in = c.execute(
            """SELECT COALESCE(SUM(py.amount), 0)
               FROM payments py JOIN payment_methods pm ON pm.id = py.method_id
               WHERE py.shift_id = %s AND py.direction = 'IN' AND pm.kind = 'CASH'""",
            (int(shift_id),),
        ).fetchone()[0]
        cash_out = c.execute(
            """SELECT COALESCE(SUM(py.amount), 0)
               FROM payments py JOIN payment_methods pm ON pm.id = py.method_id
               WHERE py.shift_id = %s AND py.direction = 'OUT' AND pm.kind = 'CASH'""",
            (int(shift_id),),
        ).fetchone()[0]
        cash_expenses = c.execute(
            """SELECT COALESCE(SUM(e.amount), 0)
               FROM expenses e JOIN payment_methods pm ON pm.id = e.method_id
               WHERE e.shift_id = %s AND pm.kind = 'CASH'""",
            (int(shift_id),),
        ).fetchone()[0]
        cash_returns = c.execute(
            """SELECT COALESCE(SUM(r.total), 0)
               FROM sales r
               JOIN sales original ON original.id = r.original_sale_id
               WHERE r.shift_id = %s AND r.status = 'POSTED' AND r.sale_type = 'RETURN'
                 AND r.customer_id IS NULL
                 AND EXISTS (
                   SELECT 1 FROM payment_allocations a
                   JOIN payments py ON py.id = a.payment_id
                   JOIN payment_methods pm ON pm.id = py.method_id
                   WHERE a.document_type = 'SALE' AND a.document_id = original.id
                     AND py.direction = 'IN' AND pm.kind = 'CASH'
                 )""",
            (int(shift_id),),
        ).fetchone()[0]
        expected = (
            int(shift[0] or 0) + int(cash_in or 0) - int(cash_out or 0)
            - int(cash_expenses or 0) - int(cash_returns or 0)
        )
        c.execute(
            """UPDATE cash_shifts SET closed_at = CURRENT_TIMESTAMP,
               closing_cash = %s, expected_cash = %s, status = 'CLOSED'
               WHERE id = %s""",
            (closing_cash_minor, expected, int(shift_id)),
        )
        _audit(c, user_id, "CLOSE_SHIFT", "cash_shifts", int(shift_id), {
            "status": "CLOSED", "closing_cash": closing_cash_minor,
            "expected_cash": expected,
        })
        return expected


# ---------------------------------------------------------------------------
# Sale posting
# ---------------------------------------------------------------------------
def create_cashier_sale(
    items: list[dict[str, Any]],
    cashier_id: int,
    warehouse_id: int,
    payment_method_id: int,
    amount_paid_minor: int,
    customer_id: int | None = None,
    new_customer_name: str | None = None,
    new_customer_phone: str | None = None,
    shift_id: int | None = None,
    note: str = "",
) -> dict[str, Any]:
    """Post sale lines, stock movements, payment allocation and journals atomically."""
    if not items:
        raise ValueError("السلة فارغة")
    if shift_id is None:
        raise ValueError("ابدأ وردية الكاشير قبل إتمام البيع حتى يمكن تسوية النقدية")
    try:
        amount_paid_minor = int(amount_paid_minor)
        warehouse_id, cashier_id, payment_method_id = (
            int(warehouse_id), int(cashier_id), int(payment_method_id)
        )
    except (TypeError, ValueError, OverflowError):
        raise ValueError("بيانات البيع غير صحيحة") from None
    if amount_paid_minor < 0:
        raise ValueError("المبلغ المدفوع لا يمكن أن يكون سالبًا")

    # Aggregate duplicate cart lines so stock availability is checked against
    # the combined quantity rather than each repeated row independently.
    quantities: dict[int, Decimal] = {}
    for item in items:
        try:
            product_id = int(item["product_id"])
            quantity = Decimal(str(item["quantity"]))
        except (KeyError, TypeError, ValueError, InvalidOperation):
            raise ValueError("بيانات المنتجات أو الكميات غير صحيحة") from None
        if product_id <= 0 or not quantity.is_finite() or quantity <= 0:
            raise ValueError("كل كمية يجب أن تكون رقمًا موجبًا")
        quantities[product_id] = quantities.get(product_id, Decimal("0")) + quantity

    with db.transaction() as c:
        warehouse = c.execute(
            "SELECT is_active FROM warehouses WHERE id = %s FOR UPDATE",
            (warehouse_id,),
        ).fetchone()
        if not warehouse or not warehouse[0]:
            raise ValueError("اختر مخزنًا نشطًا")
        _require_active_user(c, cashier_id, "حساب الكاشير غير موجود أو غير نشط")
        method = c.execute(
            """SELECT pm.account_id, pm.kind, pm.is_active
               FROM payment_methods pm
               JOIN accounts a ON a.id = pm.account_id AND a.is_active = 1
               WHERE pm.id = %s FOR UPDATE""",
            (payment_method_id,),
        ).fetchone()
        if not method or not method[2] or not method[0]:
            raise ValueError("طريقة الدفع غير نشطة أو غير مرتبطة بحساب أستاذ نشط")
        _require_open_shift(c, shift_id, cashier_id)

        new_customer_name = (new_customer_name or "").strip()
        new_customer_phone = (new_customer_phone or "").strip() or None
        if new_customer_name:
            if customer_id is not None:
                raise ValueError("اختر عميلًا موجودًا أو أدخل عميلًا جديدًا، وليس الاثنين معًا")
            existing = c.execute(
                """SELECT id FROM partners
                   WHERE is_customer = 1 AND LOWER(name) = LOWER(%s)
                     AND COALESCE(phone, '') = COALESCE(%s, '')
                   LIMIT 1 FOR UPDATE""",
                (new_customer_name, new_customer_phone),
            ).fetchone()
            if existing:
                customer_id = int(existing[0])
            else:
                c.execute(
                    """INSERT INTO partners(name, phone, is_customer, is_supplier)
                       VALUES (%s, %s, 1, 0)""",
                    (new_customer_name, new_customer_phone),
                )
                customer_id = int(c.lastrowid)
                _audit(c, cashier_id, "CREATE", "partners", customer_id, {
                    "name": new_customer_name, "phone": new_customer_phone,
                    "is_customer": True,
                }, "إنشاء عميل من صفحة الكاشير")

        if customer_id is not None:
            customer = c.execute(
                """SELECT is_customer, is_active FROM partners
                   WHERE id = %s FOR UPDATE""",
                (int(customer_id),),
            ).fetchone()
            if not customer or not customer[0] or not customer[1]:
                raise ValueError("العميل غير موجود أو غير نشط")
            customer_id = int(customer_id)

        allow_negative = _setting(c, "allow_negative_stock", "0") == "1"
        line_rows: list[dict[str, Any]] = []
        subtotal = tax_total = cogs_total = 0

        # Consistent ascending order for product row locks reduces deadlocks.
        for pid in sorted(quantities):
            quantity = quantities[pid]
            product = c.execute(
                """SELECT p.name, p.sell_price, p.cost_price, p.track_stock,
                          p.is_active, p.tax_rate_id,
                          COALESCE((SELECT tr.rate_bp FROM tax_rates tr
                                    WHERE tr.id = p.tax_rate_id AND tr.is_active = 1), 0)
                   FROM products p WHERE p.id = %s FOR UPDATE""",
                (pid,),
            ).fetchone()
            if not product or not product[4]:
                raise ValueError(f"المنتج رقم {pid} غير موجود أو غير نشط")
            name, unit_price, unit_cost, track_stock, _, _tax_rate_id, tax_bp = product
            if unit_price is None or int(unit_price) < 0:
                raise ValueError(f"سعر البيع غير صحيح للمنتج {name}")
            if unit_cost is None or int(unit_cost) < 0:
                raise ValueError(f"تكلفة المنتج غير صحيحة للمنتج {name}")
            if not quantity.is_finite() or quantity <= 0:
                raise ValueError(f"كمية غير صالحة للمنتج {name}")

            if track_stock:
                stock = c.execute(
                    """SELECT quantity FROM stock_levels
                       WHERE product_id = %s AND warehouse_id = %s FOR UPDATE""",
                    (pid, warehouse_id),
                ).fetchone()
                available = Decimal(str(stock[0])) if stock else Decimal("0")
                if quantity > available and not allow_negative:
                    raise ValueError(f"المخزون غير كافٍ للمنتج {name}: المتاح {available}")

            base = _money_round(quantity * Decimal(int(unit_price)))
            tax = _money_round(Decimal(base) * Decimal(int(tax_bp or 0)) / Decimal(10000))
            line_total = base + tax
            cogs = _money_round(quantity * Decimal(int(unit_cost))) if track_stock else 0
            subtotal += base
            tax_total += tax
            cogs_total += cogs
            line_rows.append({
                "product_id": pid, "name": name, "quantity": quantity,
                "unit_price": int(unit_price), "unit_cost": int(unit_cost),
                "tax_rate_bp": int(tax_bp or 0), "tax_amount": tax,
                "line_total": line_total, "base": base, "cogs": cogs,
                "track_stock": bool(track_stock),
            })

        total = subtotal + tax_total
        if total <= 0:
            raise ValueError("لا يمكن إتمام فاتورة بقيمة صفر")
        if amount_paid_minor > total:
            raise ValueError("المبلغ المدفوع لا يمكن أن يتجاوز إجمالي الفاتورة")
        if customer_id is None and amount_paid_minor != total:
            raise ValueError("البيع النقدي المباشر يجب دفعه بالكامل؛ اختر عميلًا لتسجيل المديونية")

        sale_number = _next_number(c, "SALE")
        c.execute(
            """INSERT INTO sales
               (number, sale_type, customer_id, warehouse_id, cashier_id, shift_id,
                status, subtotal, tax_total, discount, total, note)
               VALUES (%s, 'SALE', %s, %s, %s, %s, 'POSTED', %s, %s, 0, %s, %s)""",
            (sale_number, customer_id, warehouse_id, cashier_id, shift_id,
             subtotal, tax_total, total, note or None),
        )
        sale_id = int(c.lastrowid)

        for line in line_rows:
            c.execute(
                """INSERT INTO sale_items
                   (sale_id, product_id, product_name, quantity, unit_price, unit_cost,
                    line_discount, tax_rate_bp, tax_amount, line_total)
                   VALUES (%s, %s, %s, %s, %s, %s, 0, %s, %s, %s)""",
                (sale_id, line["product_id"], line["name"], line["quantity"],
                 line["unit_price"], line["unit_cost"], line["tax_rate_bp"],
                 line["tax_amount"], line["line_total"]),
            )
            if line["track_stock"]:
                c.execute(
                    """INSERT INTO stock_movements
                       (product_id, warehouse_id, movement_type, quantity_change,
                        unit_cost, reference_type, reference_id, user_id, note)
                       VALUES (%s, %s, 'SALE', %s, %s, 'SALE', %s, %s, %s)""",
                    (line["product_id"], warehouse_id, -line["quantity"],
                     line["unit_cost"], sale_id, cashier_id, note or None),
                )

        # Sales are initially posted to receivables; the corresponding payment
        # journal credits receivables in the same transaction when paid.
        ar_id = _configured_account_id(c, "account.receivable")
        sales_account_id = _configured_account_id(c, "account.sales")
        inventory_account_id = _configured_account_id(c, "account.inventory")
        cogs_account_id = _configured_account_id(c, "account.cogs")
        tax_account_id = _configured_account_id(c, "account.tax_payable")
        sale_lines: list[dict[str, Any]] = [
            {"account_id": ar_id, "debit": total, "credit": 0, "partner_id": customer_id},
            {"account_id": sales_account_id, "debit": 0, "credit": subtotal, "partner_id": customer_id},
        ]
        if tax_total:
            sale_lines.append({"account_id": tax_account_id, "debit": 0,
                               "credit": tax_total, "partner_id": customer_id})
        if cogs_total:
            sale_lines.extend([
                {"account_id": cogs_account_id, "debit": cogs_total, "credit": 0, "partner_id": customer_id},
                {"account_id": inventory_account_id, "debit": 0, "credit": cogs_total, "partner_id": customer_id},
            ])
        _post_journal(c, f"فاتورة بيع {sale_number}", "SALE", sale_id, cashier_id, sale_lines)

        payment_id = None
        if amount_paid_minor > 0:
            payment_number = _next_number(c, "PAYMENT")
            c.execute(
                """INSERT INTO payments
                   (number, direction, partner_id, method_id, shift_id, amount, note, user_id)
                   VALUES (%s, 'IN', %s, %s, %s, %s, %s, %s)""",
                (payment_number, customer_id, payment_method_id, shift_id,
                 amount_paid_minor, f"تحصيل فاتورة {sale_number}", cashier_id),
            )
            payment_id = int(c.lastrowid)
            c.execute(
                """INSERT INTO payment_allocations(payment_id, document_type, document_id, amount)
                   VALUES (%s, 'SALE', %s, %s)""",
                (payment_id, sale_id, amount_paid_minor),
            )
            _post_journal(
                c, f"تحصيل {payment_number} للفاتورة {sale_number}",
                "PAYMENT", payment_id, cashier_id,
                [
                    {"account_id": int(method[0]), "debit": amount_paid_minor,
                     "credit": 0, "partner_id": customer_id},
                    {"account_id": ar_id, "debit": 0,
                     "credit": amount_paid_minor, "partner_id": customer_id},
                ],
            )
            _audit(c, cashier_id, "CREATE", "payments", payment_id, {
                "number": payment_number, "direction": "IN", "sale_id": sale_id,
                "amount": amount_paid_minor, "method_id": payment_method_id,
            })

        _audit(c, cashier_id, "POST", "sales", sale_id, {
            "number": sale_number, "customer_id": customer_id,
            "warehouse_id": warehouse_id, "subtotal": subtotal,
            "tax_total": tax_total, "total": total,
            "amount_paid": amount_paid_minor, "payment_id": payment_id,
        })
        return {
            "sale_id": sale_id, "number": sale_number, "subtotal": subtotal,
            "tax_total": tax_total, "total": total, "amount_paid": amount_paid_minor,
            "remaining": total - amount_paid_minor, "payment_id": payment_id,
            "item_count": len(line_rows), "customer_id": customer_id,
        }


# ---------------------------------------------------------------------------
# Sales returns
# ---------------------------------------------------------------------------
def get_returnable_sales_items(sale_identifier: str | int) -> tuple[int, str, list[tuple]]:
    identifier = str(sale_identifier).strip()
    if not identifier:
        raise ValueError("أدخل رقم الفاتورة أو معرفها")
    numeric_id = int(identifier) if identifier.isdigit() else -1
    row = db.fetch_one(
        """SELECT id, number, status, sale_type FROM sales
           WHERE id = %s OR number = %s ORDER BY id DESC LIMIT 1""",
        (numeric_id, identifier),
    )
    if not row or row[2] != "POSTED" or row[3] != "SALE":
        raise ValueError("فاتورة البيع الأصلية غير موجودة أو غير مرحلة")

    sale_id, sale_number = int(row[0]), row[1]
    items = db.fetch_all(
        """SELECT si.product_id, si.product_name, SUM(si.quantity),
                  MIN(si.unit_price), MIN(si.unit_cost)
           FROM sale_items si WHERE si.sale_id = %s
           GROUP BY si.product_id, si.product_name ORDER BY si.product_name""",
        (sale_id,),
    )
    result = []
    for pid, name, sold_qty, unit_price, unit_cost in items:
        returned = db.fetch_one(
            """SELECT COALESCE(SUM(ri.quantity), 0)
               FROM sales r JOIN sale_items ri ON ri.sale_id = r.id
               WHERE r.original_sale_id = %s AND r.sale_type = 'RETURN'
                 AND r.status = 'POSTED' AND ri.product_id = %s""",
            (sale_id, int(pid)),
        )
        remaining = Decimal(str(sold_qty or 0)) - Decimal(str(returned[0] if returned else 0))
        if remaining > Decimal("0.000000001"):
            result.append((int(pid), name, float(remaining), int(unit_price), int(unit_cost)))
    return sale_id, sale_number, result


def create_cashier_return(
    original_sale_id: int,
    product_id: int,
    quantity: float,
    cashier_id: int,
    shift_id: int | None = None,
    note: str = "",
) -> int:
    try:
        qty = Decimal(str(quantity))
        original_sale_id, product_id, cashier_id = int(original_sale_id), int(product_id), int(cashier_id)
    except (InvalidOperation, TypeError, ValueError, OverflowError):
        raise ValueError("بيانات المرتجع غير صحيحة") from None
    if not qty.is_finite() or qty <= 0:
        raise ValueError("كمية المرتجع يجب أن تكون موجبة")
    if shift_id is None:
        raise ValueError("ابدأ وردية الكاشير قبل تسجيل المرتجع")

    with db.transaction() as c:
        _require_active_user(c, cashier_id, "حساب الكاشير غير موجود أو غير نشط")
        _require_open_shift(c, shift_id, cashier_id)
        original = c.execute(
            """SELECT number, customer_id, warehouse_id, status, sale_type
               FROM sales WHERE id = %s FOR UPDATE""",
            (original_sale_id,),
        ).fetchone()
        if not original or original[3] != "POSTED" or original[4] != "SALE":
            raise ValueError("فاتورة البيع الأصلية غير موجودة أو غير مرحلة")

        line = c.execute(
            """SELECT product_name, SUM(quantity), MIN(unit_price), MIN(unit_cost),
                      SUM(line_discount), MAX(tax_rate_bp)
               FROM sale_items WHERE sale_id = %s AND product_id = %s
               GROUP BY product_id, product_name""",
            (original_sale_id, product_id),
        ).fetchone()
        if not line:
            raise ValueError("المنتج غير موجود ضمن فاتورة البيع الأصلية")
        product_name, sold_qty_raw, unit_price, unit_cost, total_discount, tax_rate_bp = line
        sold_qty = Decimal(str(sold_qty_raw or 0))
        returned = c.execute(
            """SELECT COALESCE(SUM(ri.quantity), 0)
               FROM sales r JOIN sale_items ri ON ri.sale_id = r.id
               WHERE r.original_sale_id = %s AND r.sale_type = 'RETURN'
                 AND r.status = 'POSTED' AND ri.product_id = %s""",
            (original_sale_id, product_id),
        ).fetchone()
        remaining = sold_qty - Decimal(str(returned[0] or 0))
        if qty > remaining:
            raise ValueError(f"الكمية المتبقية للمرتجع {max(remaining, Decimal('0')):g} فقط")

        product = c.execute(
            "SELECT track_stock FROM products WHERE id = %s FOR UPDATE", (product_id,)
        ).fetchone()
        if product is None:
            raise ValueError("المنتج غير موجود")
        track_stock = bool(product[0])

        base = _money_round(qty * Decimal(int(unit_price)))
        discount = (
            min(base, _money_round(Decimal(int(total_discount or 0)) * qty / sold_qty))
            if sold_qty else 0
        )
        tax_bp = int(tax_rate_bp or 0)
        tax_amount = _money_round(Decimal(base - discount) * Decimal(tax_bp) / Decimal(10000))
        total = base - discount + tax_amount
        if total <= 0:
            raise ValueError("لا يمكن إنشاء مرتجع بقيمة صفر")

        number = _next_number(c, "SALE_RETURN")
        c.execute(
            """INSERT INTO sales
               (number, sale_type, original_sale_id, customer_id, warehouse_id,
                cashier_id, shift_id, status, subtotal, tax_total, discount, total, note)
               VALUES (%s, 'RETURN', %s, %s, %s, %s, %s, 'POSTED', %s, %s, %s, %s, %s)""",
            (number, original_sale_id, original[1], original[2], cashier_id, shift_id,
             base, tax_amount, discount, total, note or f"مرتجع من {original[0]}"),
        )
        return_id = int(c.lastrowid)
        c.execute(
            """INSERT INTO sale_items
               (sale_id, product_id, product_name, quantity, unit_price, unit_cost,
                line_discount, tax_rate_bp, tax_amount, line_total)
               VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)""",
            (return_id, product_id, product_name, qty, int(unit_price), int(unit_cost),
             discount, tax_bp, tax_amount, total),
        )

        if track_stock:
            c.execute(
                """INSERT INTO stock_movements
                   (product_id, warehouse_id, movement_type, quantity_change, unit_cost,
                    reference_type, reference_id, user_id, note)
                   VALUES (%s, %s, 'SALE_RETURN', %s, %s, 'SALE_RETURN', %s, %s, %s)""",
                (product_id, int(original[2]), qty, int(unit_cost), return_id,
                 cashier_id, note or None),
            )

        returns_account = _configured_account_id(c, "account.sales_returns")
        tax_account = _configured_account_id(c, "account.tax_payable")
        ar_account = _configured_account_id(c, "account.receivable")
        lines: list[dict[str, Any]] = [{
            "account_id": returns_account, "debit": base - discount,
            "credit": 0, "partner_id": original[1],
        }]
        if tax_amount:
            lines.append({"account_id": tax_account, "debit": tax_amount,
                          "credit": 0, "partner_id": original[1]})

        # Walk-in returns credit the method's ledger account used on the original
        # receipt where possible; customer-account returns credit receivables.
        if original[1] is not None:
            counter_account = ar_account
        else:
            receipt = c.execute(
                """SELECT pm.account_id FROM payment_allocations a
                   JOIN payments py ON py.id = a.payment_id
                   JOIN payment_methods pm ON pm.id = py.method_id
                   JOIN accounts ac ON ac.id = pm.account_id AND ac.is_active = 1
                   WHERE a.document_type = 'SALE' AND a.document_id = %s
                     AND py.direction = 'IN'
                   ORDER BY py.id LIMIT 1""",
                (original_sale_id,),
            ).fetchone()
            counter_account = int(receipt[0]) if receipt and receipt[0] else _configured_account_id(c, "account.cash")
        lines.append({"account_id": counter_account, "debit": 0,
                      "credit": total, "partner_id": original[1]})

        if track_stock:
            cogs_reversal = _money_round(qty * Decimal(int(unit_cost)))
            if cogs_reversal:
                lines.extend([
                    {"account_id": _configured_account_id(c, "account.inventory"),
                     "debit": cogs_reversal, "credit": 0, "partner_id": original[1]},
                    {"account_id": _configured_account_id(c, "account.cogs"),
                     "debit": 0, "credit": cogs_reversal, "partner_id": original[1]},
                ])
        _post_journal(c, f"مرتجع بيع {number}", "SALE_RETURN", return_id, cashier_id, lines)
        _audit(c, cashier_id, "POST_RETURN", "sales", return_id, {
            "number": number, "original_sale_id": original_sale_id,
            "product_id": product_id, "quantity": str(qty), "total": total,
        })
        return return_id

# ---------------------------------------------------------------------------
# Excel sales register (called only after the sale transaction commits)
# ---------------------------------------------------------------------------
def export_cashier_sale_to_excel(sale_id: int, shift_id: int | None = None) -> str:
    """Create/update the daily Excel sales register without changing the ERP DB."""
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("تصدير Excel يحتاج تثبيت المكتبة: pip install openpyxl") from exc

    sale_id = int(sale_id)
    sale = db.fetch_one(
        """SELECT id, number, sale_type, customer_id, warehouse_id, cashier_id,
                  shift_id, status, subtotal, tax_total, discount, total, note
           FROM sales WHERE id = %s""",
        (sale_id,),
    )
    if not sale:
        raise ValueError("لم يتم العثور على الفاتورة المطلوبة لتصديرها")
    if sale[7] != "POSTED" or sale[2] != "SALE":
        raise ValueError("تصدير Excel متاح لفواتير البيع المرحلة فقط")

    sale_shift_id = int(sale[6]) if sale[6] is not None else None
    if sale_shift_id is None:
        raise ValueError("لا توجد وردية مرتبطة بالفاتورة")
    if shift_id is not None and int(shift_id) != sale_shift_id:
        raise ValueError("رقم الوردية لا يطابق الوردية المسجلة على الفاتورة")

    folder = Path(get_cashier_session_folder(sale_shift_id, create=True))
    output_path = folder / "Sales_Register.xlsx"
    labels = db.fetch_one(
        """SELECT COALESCE(p.name, 'عميل نقدي'), COALESCE(p.phone, ''),
                  COALESCE(w.name, ''), COALESCE(u.username, '')
           FROM sales s
           LEFT JOIN partners p ON p.id = s.customer_id
           LEFT JOIN warehouses w ON w.id = s.warehouse_id
           LEFT JOIN users u ON u.id = s.cashier_id
           WHERE s.id = %s""",
        (sale_id,),
    )
    customer_name, customer_phone, warehouse_name, cashier_name = labels or ("عميل نقدي", "", "", "")
    shift_row = db.fetch_one(
        "SELECT opened_at FROM cash_shifts WHERE id = %s", (sale_shift_id,)
    )
    session_started = str(shift_row[0]) if shift_row and shift_row[0] else date.today().isoformat()

    items = db.fetch_all(
        """SELECT si.product_id, si.product_name, si.quantity, si.unit_price,
                  si.line_discount, si.tax_rate_bp, si.tax_amount, si.line_total,
                  COALESCE(p.sku, ''),
                  COALESCE((SELECT pb.barcode FROM product_barcodes pb
                            WHERE pb.product_id = si.product_id
                            ORDER BY pb.is_primary DESC, pb.id LIMIT 1), '')
           FROM sale_items si LEFT JOIN products p ON p.id = si.product_id
           WHERE si.sale_id = %s ORDER BY si.id""",
        (sale_id,),
    )
    payments = db.fetch_all(
        """SELECT py.number, pm.name, py.amount, a.amount, py.direction
           FROM payment_allocations a
           JOIN payments py ON py.id = a.payment_id
           JOIN payment_methods pm ON pm.id = py.method_id
           WHERE a.document_type = 'SALE' AND a.document_id = %s
           ORDER BY py.id""",
        (sale_id,),
    )

    decimals = get_currency_decimals()
    factor = Decimal(10) ** decimals
    money_format = "#,##0" + ("." + ("0" * decimals) if decimals else "")

    def major(minor: Any) -> float:
        return float(Decimal(int(minor or 0)) / factor)

    def qty_text(value: Any) -> str:
        return format(Decimal(str(value or 0)).normalize(), "f")

    def safe_text(value: Any) -> str:
        """Prevent user-controlled text from being evaluated as an Excel formula."""
        result = "" if value is None else str(value)
        if result.lstrip().startswith(("=", "+", "-", "@")):
            return "'" + result
        return result

    item_lines: list[str] = []
    for item in items:
        pid, name, qty, unit_price, discount, _tax_bp, tax_amount, line_total, sku, barcode = item
        identifiers = [f"ID {pid}"]
        if sku:
            identifiers.append(f"SKU {safe_text(sku)}")
        if barcode:
            identifiers.append(f"باركود {safe_text(barcode)}")
        item_lines.append(
            f"{safe_text(name)} ({' | '.join(identifiers)}) — "
            f"الكمية: {qty_text(qty)}؛ سعر الوحدة: {major(unit_price):.{decimals}f}؛ "
            f"الخصم: {major(discount):.{decimals}f}؛ "
            f"الضريبة: {major(tax_amount):.{decimals}f}؛ "
            f"إجمالي السطر: {major(line_total):.{decimals}f}"
        )
    item_details = safe_text("\n".join(item_lines) if item_lines else "لا توجد تفاصيل أصناف")
    incoming_paid = sum(
        int(payment[3] or 0)
        for payment in payments
        if str(payment[4]).upper() == "IN"
    )
    payment_details = safe_text("\n".join(
        f"{safe_text(payment[0])} — {safe_text(payment[1])}: {major(payment[3]):.{decimals}f}"
        for payment in payments
    ) or "—")
    remaining = max(0, int(sale[11] or 0) - incoming_paid)

    headers = [
        "رقم العملية / الفاتورة (للمرتجع)", "تاريخ بدء الوردية", "رقم الوردية",
        "الكاشير", "العميل", "هاتف العميل", "المخزن", "تفاصيل الأصناف المباعة",
        "عدد بنود الأصناف", "الإجمالي قبل الضريبة", "الخصم", "الضريبة",
        "إجمالي الفاتورة", "المدفوع", "المتبقي", "الإيصالات وطرق الدفع", "ملاحظات",
    ]
    row_values = [
        safe_text(sale[1]), session_started, sale_shift_id, safe_text(cashier_name),
        safe_text(customer_name), safe_text(customer_phone), safe_text(warehouse_name),
        item_details, len(items), major(sale[8]), major(sale[10]), major(sale[9]),
        major(sale[11]), major(incoming_paid), major(remaining), payment_details,
        safe_text(sale[12] or "—"),
    ]

    wb = None
    if output_path.exists():
        try:
            wb = load_workbook(output_path)
        except Exception as exc:
            raise RuntimeError(
                "تعذر فتح سجل المبيعات الحالي. أغلق ملف Excel إن كان مفتوحًا ثم أعد المحاولة."
            ) from exc
    else:
        wb = Workbook()

    try:
        # Reuse a sheet whose register header is valid, even if an earlier
        # unrelated worksheet forced the register to be named "سجل المبيعات 2".
        ws = next(
            (
                wb[name]
                for name in wb.sheetnames
                if wb[name].max_row >= 3
                and wb[name].cell(row=3, column=1).value == headers[0]
            ),
            None,
        )
        if ws is None:
            default_sheet_is_empty = (
                len(wb.sheetnames) == 1
                and wb.active.max_row == 1
                and wb.active.cell(1, 1).value is None
            )
            if default_sheet_is_empty:
                ws = wb.active
                ws.title = "سجل المبيعات"
            else:
                base_name = "سجل المبيعات"
                sheet_name = base_name
                suffix = 2
                while sheet_name in wb.sheetnames:
                    sheet_name = f"{base_name} {suffix}"
                    suffix += 1
                ws = wb.create_sheet(sheet_name)

        ws.sheet_view.rightToLeft = True
        ws.freeze_panes = "A4"
        navy, pale, white, dark = "17365D", "F3F6FA", "FFFFFF", "263238"
        thin_gray = Side(style="thin", color="D7DEE8")
        border = Border(bottom=thin_gray)
        title_font = Font(name="Arial", size=16, bold=True, color=white)
        header_font = Font(name="Arial", size=10, bold=True, color=white)
        body_font = Font(name="Arial", size=10, color=dark)

        header_is_valid = ws.max_row >= 3 and ws.cell(row=3, column=1).value == headers[0]
        if not header_is_valid:
            ws.sheet_view.rightToLeft = True
            ws.freeze_panes = "A4"
            ws.merge_cells("A1:Q1")
            ws["A1"] = "سجل عمليات البيع — MyShop ERP"
            ws["A1"].fill = PatternFill("solid", fgColor=navy)
            ws["A1"].font = title_font
            ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[1].height = 30
            ws.merge_cells("A2:Q2")
            ws["A2"] = (
                "لإرجاع منتج: انسخ رقم العملية من العمود الأول وأدخله في نافذة مرتجع بيع. "
                "يُسجل كل بيع في صف واحد وتظهر تفاصيل الأصناف داخل عمودها."
            )
            ws["A2"].font = Font(name="Arial", size=10, italic=True, color=dark)
            ws["A2"].alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
            ws.row_dimensions[2].height = 30
            for col_index, heading in enumerate(headers, start=1):
                cell = ws.cell(row=3, column=col_index, value=heading)
                cell.fill = PatternFill("solid", fgColor=navy)
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = border
            ws.row_dimensions[3].height = 36

        target_row = None
        for row_index in range(4, ws.max_row + 1):
            if str(ws.cell(row=row_index, column=1).value or "") == str(sale[1]):
                target_row = row_index
                break
        if target_row is None:
            target_row = max(4, ws.max_row + 1)

        for col_index, value in enumerate(row_values, start=1):
            cell = ws.cell(row=target_row, column=col_index, value=value)
            cell.font = body_font
            cell.alignment = Alignment(
                horizontal="right" if col_index in (4, 5, 6, 8, 16, 17) else "center",
                vertical="top", wrap_text=True,
            )
            cell.border = border
            if target_row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=pale)
            if col_index in (10, 11, 12, 13, 14, 15):
                cell.number_format = money_format

        wrapped_lines = max(1, len(item_lines), len(payment_details.splitlines()))
        ws.row_dimensions[target_row].height = min(180, max(34, wrapped_lines * 28))
        widths = [27, 22, 12, 18, 24, 17, 22, 68, 15, 19, 15, 15, 19, 17, 17, 38, 32]
        for col_index, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(col_index)].width = width
        ws.auto_filter.ref = f"A3:Q{max(3, ws.max_row)}"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.outlinePr.summaryBelow = True

        fd, temp_name = tempfile.mkstemp(
            prefix=".sales_register_", suffix=".xlsx", dir=str(folder)
        )
        os.close(fd)
        try:
            wb.save(temp_name)
            os.replace(temp_name, output_path)
        except PermissionError as exc:
            raise RuntimeError(
                "تعذر تحديث سجل المبيعات. أغلق Sales_Register.xlsx في Excel ثم أعد تصدير العملية."
            ) from exc
        finally:
            try:
                if os.path.exists(temp_name):
                    os.unlink(temp_name)
            except OSError:
                pass
    finally:
        if wb is not None:
            wb.close()
    return str(output_path)

# ---------------------------------------------------------------------------
# Customer debt operations
# Debt = customer balance view + append-only manual debt adjustments.
# Every mutation is transactional and writes a journal entry plus audit event.
# ---------------------------------------------------------------------------
DEBT_JOURNAL_SOURCE = "DEBT_ADJUSTMENT"
NOT_CUSTOMER_MSG = (
    "هذا الشخص ليس عميلًا مسجلًا؛ لا يُسمح بالدين إلا للعملاء المسجلين، "
    "ويجب دفع ثمن المنتجات كاملًا."
)


def _require_debtor_customer(c, customer_id: int | None, lock: bool = False) -> str:
    if customer_id is None:
        raise ValueError(NOT_CUSTOMER_MSG)
    sql = "SELECT name, is_customer, is_active FROM partners WHERE id = %s"
    if lock:
        sql += " FOR UPDATE"
    row = c.execute(sql, (int(customer_id),)).fetchone()
    if not row or not row[1]:
        raise ValueError(NOT_CUSTOMER_MSG)
    if not row[2]:
        raise ValueError("العميل غير نشط")
    return str(row[0])


def _customer_debt(c, customer_id: int) -> int:
    row = c.execute(
        "SELECT COALESCE(balance, 0) FROM v_customer_balances WHERE partner_id = %s",
        (int(customer_id),),
    ).fetchone()
    base = int(row[0]) if row and row[0] is not None else 0
    adjustment = c.execute(
        """SELECT COALESCE(SUM(CASE kind WHEN 'ADD' THEN amount ELSE -amount END), 0)
           FROM customer_debt_adjustments WHERE partner_id = %s""",
        (int(customer_id),),
    ).fetchone()[0]
    return base + int(adjustment or 0)


def _first_configured_account(c, *keys: str) -> int:
    for key in keys:
        code = _setting(c, key)
        if code in (None, ""):
            continue
        row = c.execute(
            "SELECT id FROM accounts WHERE code = %s AND is_active = 1 LIMIT 1",
            (str(code),),
        ).fetchone()
        if row:
            return int(row[0])
    raise ValueError("إعداد حساب تسوية الديون غير موجود: " + " / ".join(keys))


def _clean_reason(reason: str | None) -> str:
    text = (reason or "").strip()
    if not text:
        raise ValueError("اكتب سببًا لتعديل الدين (يظهر في سجل المدير)")
    return text


def _positive_minor(value: Any, message: str) -> int:
    try:
        amount = int(value)
    except (TypeError, ValueError, OverflowError):
        raise ValueError(message) from None
    if amount <= 0:
        raise ValueError(message)
    return amount


def _apply_debt_adjustment(
    c,
    customer_id: int,
    customer_name: str,
    kind: str,
    amount: int,
    reason: str,
    user_id: int | None,
    shift_id: int | None,
    audit_action: str,
    debt_before: int,
) -> dict[str, Any]:
    if kind not in {"ADD", "REMOVE"}:
        raise ValueError("نوع تسوية الدين غير صالح")
    number = _next_number(c, "DEBT_ADJ")
    c.execute(
        """INSERT INTO customer_debt_adjustments
           (number, partner_id, kind, amount, reason, user_id, shift_id)
           VALUES (%s, %s, %s, %s, %s, %s, %s)""",
        (number, int(customer_id), kind, int(amount), reason, user_id, shift_id),
    )
    adjustment_id = int(c.lastrowid)
    receivable_id = _configured_account_id(c, "account.receivable")

    if kind == "ADD":
        offset_id = _first_configured_account(c, "account.debt_adjustment", "account.sales")
        lines = [
            {"account_id": receivable_id, "debit": amount, "credit": 0, "partner_id": customer_id},
            {"account_id": offset_id, "debit": 0, "credit": amount, "partner_id": customer_id},
        ]
        debt_after = debt_before + amount
    else:
        offset_id = _first_configured_account(
            c, "account.bad_debt", "account.debt_adjustment", "account.sales_returns"
        )
        lines = [
            {"account_id": offset_id, "debit": amount, "credit": 0, "partner_id": customer_id},
            {"account_id": receivable_id, "debit": 0, "credit": amount, "partner_id": customer_id},
        ]
        debt_after = debt_before - amount

    _post_journal(
        c, f"تسوية دين العميل {customer_name} ({number})",
        DEBT_JOURNAL_SOURCE, adjustment_id, user_id, lines,
    )
    _audit(c, user_id, audit_action, "customer_debt_adjustments", adjustment_id, {
        "number": number, "customer_id": int(customer_id), "customer": customer_name,
        "kind": kind, "amount": amount,
        "debt_before": debt_before, "debt_after": debt_after,
    }, reason)
    return {
        "adjustment_id": adjustment_id, "number": number, "kind": kind,
        "amount": amount, "debt_before": debt_before, "debt_after": debt_after,
    }


def get_debt_customers() -> list[tuple[int, str, str, int]]:
    """Return active customers with positive outstanding debt."""
    rows = db.fetch_all(
        """SELECT id, name, phone, debt FROM (
               SELECT p.id AS id, p.name AS name, COALESCE(p.phone, '') AS phone,
                      COALESCE(b.balance, 0) + COALESCE(adj.net, 0) AS debt
               FROM partners p
               LEFT JOIN v_customer_balances b ON b.partner_id = p.id
               LEFT JOIN (
                   SELECT partner_id,
                          SUM(CASE kind WHEN 'ADD' THEN amount ELSE -amount END) AS net
                   FROM customer_debt_adjustments GROUP BY partner_id
               ) adj ON adj.partner_id = p.id
               WHERE p.is_customer = 1 AND p.is_active = 1
           ) AS debt_rows
           WHERE debt > 0 ORDER BY debt DESC, name"""
    )
    return [(int(row[0]), str(row[1]), str(row[2]), int(row[3])) for row in rows]


def get_customer_debt(customer_id: int) -> int:
    with db.transaction() as c:
        _require_debtor_customer(c, customer_id)
        return _customer_debt(c, customer_id)


def pay_customer_debt(
    customer_id: int,
    payment_method_id: int,
    user_id: int,
    shift_id: int | None,
    amount_minor: int | None = None,
) -> dict[str, Any]:
    """Collect some/all of a customer's debt, allocating to oldest unpaid sales first."""
    if shift_id is None:
        raise ValueError("ابدأ وردية الكاشير قبل تسديد الدين")

    with db.transaction() as c:
        name = _require_debtor_customer(c, customer_id, lock=True)
        _require_active_user(c, user_id)
        _require_open_shift(c, shift_id, user_id)
        method = c.execute(
            """SELECT pm.account_id, pm.is_active
               FROM payment_methods pm
               JOIN accounts a ON a.id = pm.account_id AND a.is_active = 1
               WHERE pm.id = %s FOR UPDATE""",
            (int(payment_method_id),),
        ).fetchone()
        if not method or not method[1] or not method[0]:
            raise ValueError("طريقة الدفع غير نشطة أو غير مرتبطة بحساب أستاذ نشط")

        debt = _customer_debt(c, customer_id)
        if debt <= 0:
            raise ValueError("لا يوجد دين على هذا العميل")
        amount = debt if amount_minor is None else _positive_minor(
            amount_minor, "مبلغ التسديد يجب أن يكون أكبر من صفر"
        )
        if amount > debt:
            raise ValueError("مبلغ التسديد أكبر من الدين المستحق")

        number = _next_number(c, "PAYMENT")
        c.execute(
            """INSERT INTO payments
               (number, direction, partner_id, method_id, shift_id, amount, note, user_id)
               VALUES (%s, 'IN', %s, %s, %s, %s, %s, %s)""",
            (number, int(customer_id), int(payment_method_id), int(shift_id), amount,
             f"تسديد دين العميل {name}", int(user_id)),
        )
        payment_id = int(c.lastrowid)

        unpaid = c.execute(
            """SELECT id, outstanding FROM (
                   SELECT s.id AS id,
                          s.total - COALESCE((
                              SELECT SUM(a.amount) FROM payment_allocations a
                              WHERE a.document_type = 'SALE' AND a.document_id = s.id
                          ), 0) AS outstanding
                   FROM sales s
                   WHERE s.customer_id = %s AND s.sale_type = 'SALE' AND s.status = 'POSTED'
               ) AS unpaid_sales
               WHERE outstanding > 0 ORDER BY id FOR UPDATE""",
            (int(customer_id),),
        ).fetchall()
        remaining = amount
        allocated_sales: list[int] = []
        for sale_id, outstanding in unpaid:
            if remaining <= 0:
                break
            part = min(remaining, int(outstanding))
            if part <= 0:
                continue
            c.execute(
                """INSERT INTO payment_allocations(payment_id, document_type, document_id, amount)
                   VALUES (%s, 'SALE', %s, %s)""",
                (payment_id, int(sale_id), part),
            )
            allocated_sales.append(int(sale_id))
            remaining -= part

        receivable_id = _configured_account_id(c, "account.receivable")
        _post_journal(c, f"تسديد دين العميل {name} ({number})", "PAYMENT", payment_id, user_id, [
            {"account_id": int(method[0]), "debit": amount, "credit": 0, "partner_id": customer_id},
            {"account_id": receivable_id, "debit": 0, "credit": amount, "partner_id": customer_id},
        ])
        _audit(c, user_id, "DEBT_PAY", "payments", payment_id, {
            "number": number, "customer_id": int(customer_id), "customer": name,
            "amount": amount, "method_id": int(payment_method_id), "shift_id": int(shift_id),
            "debt_before": debt, "debt_after": debt - amount,
            "allocated_sales": allocated_sales, "on_account": remaining,
        }, "تسديد دين عميل")
        return {
            "payment_id": payment_id, "number": number, "amount": amount,
            "debt_before": debt, "remaining_debt": debt - amount,
        }


def update_customer_debt(
    customer_id: int,
    new_debt_minor: int,
    user_id: int,
    reason: str,
    shift_id: int | None = None,
) -> dict[str, Any]:
    """Set current debt via an auditable ADD or REMOVE adjustment."""
    reason = _clean_reason(reason)
    try:
        target = int(new_debt_minor)
    except (TypeError, ValueError, OverflowError):
        raise ValueError("قيمة الدين الجديدة غير صحيحة") from None
    if target < 0:
        raise ValueError("قيمة الدين الجديدة لا يمكن أن تكون سالبة")

    with db.transaction() as c:
        name = _require_debtor_customer(c, customer_id, lock=True)
        current = _customer_debt(c, customer_id)
        delta = target - current
        if delta == 0:
            raise ValueError("قيمة الدين الجديدة مطابقة للدين الحالي")
        return _apply_debt_adjustment(
            c, int(customer_id), name, "ADD" if delta > 0 else "REMOVE",
            abs(delta), reason, user_id, shift_id, "DEBT_UPDATE", current,
        )


def remove_customer_debt(
    customer_id: int,
    user_id: int,
    reason: str,
    amount_minor: int | None = None,
    shift_id: int | None = None,
) -> dict[str, Any]:
    """Write off part/all of debt without receiving money."""
    reason = _clean_reason(reason)
    with db.transaction() as c:
        name = _require_debtor_customer(c, customer_id, lock=True)
        current = _customer_debt(c, customer_id)
        if current <= 0:
            raise ValueError("لا يوجد دين على هذا العميل")
        amount = current if amount_minor is None else _positive_minor(
            amount_minor, "المبلغ المراد حذفه يجب أن يكون أكبر من صفر"
        )
        if amount > current:
            raise ValueError("المبلغ أكبر من الدين الحالي")
        return _apply_debt_adjustment(
            c, int(customer_id), name, "REMOVE", amount, reason, user_id,
            shift_id, "DEBT_REMOVE", current,
        )


def add_customer_debt(
    customer_id: int,
    amount_minor: int,
    user_id: int,
    reason: str,
    shift_id: int | None = None,
) -> dict[str, Any]:
    """Add an auditable debt adjustment to an active registered customer."""
    reason = _clean_reason(reason)
    amount = _positive_minor(amount_minor, "مبلغ الدين يجب أن يكون أكبر من صفر")
    with db.transaction() as c:
        name = _require_debtor_customer(c, customer_id, lock=True)
        current = _customer_debt(c, customer_id)
        return _apply_debt_adjustment(
            c, int(customer_id), name, "ADD", amount, reason, user_id,
            shift_id, "DEBT_ADD", current,
        )


def get_debt_log(limit: int = 200) -> list[tuple]:
    """Return latest debt audit events for the admin page."""
    try:
        bounded_limit = max(1, min(2000, int(limit)))
    except (TypeError, ValueError, OverflowError):
        bounded_limit = 200
    return db.fetch_all(
        """SELECT a.id, COALESCE(u.username, ''), a.action, a.entity,
                  a.entity_id, a.new_values, COALESCE(a.note, '')
           FROM audit_log a LEFT JOIN users u ON u.id = a.user_id
           WHERE a.action LIKE 'DEBT%'
           ORDER BY a.id DESC LIMIT %s""",
        (bounded_limit,),
    )
